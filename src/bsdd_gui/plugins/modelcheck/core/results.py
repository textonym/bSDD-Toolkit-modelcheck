import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from bsdd_gui.plugins.modelcheck.module import constants


def create_excel_report(issues: list[dict], export_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Modelcheck Issues"

    headers = constants.ISSUE_TABLE_HEADER
    ws.append(headers)

    for issue in issues:
        row = [
            issue.get("GUID_ZWC", ""),
            issue.get("GUID", ""),
            issue.get("creation_date", ""),
            issue.get("short_description", ""),
            issue.get("issue_type", ""),
            issue.get("PropertySet", ""),
            issue.get("Property", ""),
            str(issue.get("Value", "")),
            str(issue.get("ValueType", "")),
            issue.get("description", ""),
        ]
        ws.append(row)

    row_count = max(len(issues) + 1, 2)
    col_count = len(headers)
    end_col = get_column_letter(col_count)
    table_ref = f"A1:{end_col}{row_count}"

    tab = Table(displayName="ModelcheckTable", ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=True,
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(export_path)
